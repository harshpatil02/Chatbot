from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, List

from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from app.database import add_document_chunk, create_tables


def load_pdf_text(pdf_path: str) -> str:
    reader = PdfReader(pdf_path)
    texts = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(texts).strip()


def load_xlsx_text(xlsx_path: str) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    parts: List[str] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        for row in rows:
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                parts.append(" | ".join(values))
    workbook.close()
    return "\n".join(parts).strip()


def extract_text_from_file(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf_text(str(path))
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return load_xlsx_text(str(path))
    raise ValueError(f"Unsupported file type: {suffix}")


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 120) -> List[str]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=overlap)
    return [chunk for chunk in splitter.split_text(text) if chunk.strip()]


def embed_text(text: str, dimension: int = 1536) -> List[float]:
    if not text.strip():
        return [0.0] * dimension

    tokens = re.findall(r"\w+", text.lower())
    vector = [0.0] * dimension
    for token in tokens:
        index = sum(ord(char) for char in token) % dimension
        vector[index] += 1.0

    norm = sum(value * value for value in vector) ** 0.5
    if norm == 0:
        return [0.0] * dimension
    return [value / norm for value in vector]


def store_chunks(chunks: List[str], metadata: Dict[str, Any] | None = None) -> None:
    create_tables()
    metadata = metadata or {}
    for chunk in chunks:
        embedding = embed_text(chunk)
        add_document_chunk(chunk, metadata=metadata, embedding=embedding)


def ingest_file(file_path: str, metadata: Dict[str, Any] | None = None) -> int:
    text = extract_text_from_file(file_path)
    chunks = chunk_text(text)
    store_chunks(chunks, metadata=metadata)
    return len(chunks)


def ingest_pdf(pdf_path: str, metadata: Dict[str, Any] | None = None) -> int:
    return ingest_file(pdf_path, metadata=metadata)
